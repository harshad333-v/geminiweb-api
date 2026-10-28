"""Document extraction helpers with optional format-specific dependencies."""
import csv
import io
import json
import zipfile
import xml.etree.ElementTree as ET


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".pdf", ".docx", ".xlsx"}


class DocumentError(ValueError):
    """Raised when a document cannot be safely parsed."""


def extract_document(filename, data, max_bytes=10 * 1024 * 1024):
    """Extract text/records from a supported file without executing embedded content."""
    import os
    if not isinstance(data, bytes):
        raise DocumentError("file content must be bytes")
    if len(data) > max_bytes:
        raise DocumentError("file exceeds configured upload limit")
    ext = os.path.splitext(filename or "")[1].lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise DocumentError("unsupported file type; supported: " + ", ".join(sorted(SUPPORTED_EXTENSIONS)))
    if ext in (".txt", ".md"):
        return {"filename": filename, "format": ext[1:], "text": data.decode("utf-8-sig", errors="replace")}
    if ext == ".json":
        try:
            obj = json.loads(data.decode("utf-8-sig"))
        except (UnicodeDecodeError, json.JSONDecodeError) as exc:
            raise DocumentError("invalid JSON document") from exc
        return {"filename": filename, "format": "json", "data": obj}
    if ext == ".csv":
        try:
            text = data.decode("utf-8-sig")
            rows = list(csv.DictReader(io.StringIO(text)))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise DocumentError("invalid CSV document") from exc
        return {"filename": filename, "format": "csv", "rows": rows, "row_count": len(rows)}
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise DocumentError("PDF support requires the optional 'pypdf' dependency") from exc
        try:
            reader = PdfReader(io.BytesIO(data), strict=True)
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception as exc:
            raise DocumentError("unable to parse PDF document") from exc
        return {"filename": filename, "format": "pdf", "page_count": len(pages),
                "text": "\n\n".join(pages), "pages": pages}
    if ext == ".docx":
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                info = archive.getinfo("word/document.xml")
                if info.file_size > max_bytes:
                    raise DocumentError("expanded DOCX XML exceeds configured size limit")
                raw = archive.read("word/document.xml")
            root = ET.fromstring(raw)
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            paragraphs = []
            for paragraph in root.findall(".//w:p", ns):
                paragraphs.append("".join(node.text or "" for node in paragraph.findall(".//w:t", ns)))
            return {"filename": filename, "format": "docx", "text": "\n".join(paragraphs)}
        except Exception as exc:
            raise DocumentError("unable to parse DOCX document") from exc
    if ext == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise DocumentError("XLSX support requires the optional 'openpyxl' dependency") from exc
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheets = {}
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                values = [[cell for cell in row] for row in rows]
                sheets[sheet.title] = values
            workbook.close()
            return {"filename": filename, "format": "xlsx", "sheets": sheets}
        except Exception as exc:
            raise DocumentError("unable to parse XLSX document") from exc
    raise DocumentError("unsupported file type")
